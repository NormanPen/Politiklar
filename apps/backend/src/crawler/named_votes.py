"""Parser and importer for official Bundestag named-vote XLSX lists."""

from dataclasses import dataclass
from datetime import datetime
from hashlib import sha256
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from db.relational.models import NamedVote, NamedVoteRow, SourceDocument

from .fetcher import FetchedDocument, fetch_document


OUTCOME_COLUMNS = {
    "ja": "yes",
    "nein": "no",
    "enthaltung": "abstained",
    "ungültig": "invalid",
    "nichtabgegeben": "not_voted",
}


@dataclass(frozen=True)
class VoteRow:
    parliamentary_group: str | None
    last_name: str | None
    first_name: str | None
    title: str | None
    display_name: str | None
    raw_outcome: str
    outcome: str
    remark: str | None


@dataclass(frozen=True)
class ParsedNamedVote:
    electoral_term: int
    sitting_number: int
    vote_number: int
    rows: tuple[VoteRow, ...]


def parse_named_vote_xlsx(content: bytes) -> ParsedNamedVote:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    headers = tuple(_normalize_header(value) for value in next(rows))
    required = {"wahlperiode", "sitzungnr", "abstimmnr"}
    if not required.issubset(headers):
        raise ValueError("The XLSX file does not contain the required Bundestag vote columns.")

    vote_rows = [dict(zip(headers, row, strict=False)) for row in rows]
    metadata = vote_rows[0] if vote_rows else None
    if metadata is None:
        raise ValueError("The XLSX file contains no vote rows.")
    return ParsedNamedVote(
        electoral_term=_required_int(metadata, "wahlperiode"),
        sitting_number=_required_int(metadata, "sitzungnr"),
        vote_number=_required_int(metadata, "abstimmnr"),
        rows=tuple(_parse_row(row) for row in vote_rows),
    )


def import_named_vote(session: Session, url: str, title: str | None = None) -> ParsedNamedVote:
    document = fetch_document(url)
    parsed_vote = parse_named_vote_xlsx(document.content)
    source = _get_or_create_source(session, url, document)
    vote = session.scalar(select(NamedVote).where(
        NamedVote.electoral_term == parsed_vote.electoral_term,
        NamedVote.sitting_number == parsed_vote.sitting_number,
        NamedVote.vote_number == parsed_vote.vote_number,
        NamedVote.content_sha256 == document.source.content_sha256,
    ))
    if vote is None:
        vote = NamedVote(
            source_document_id=source.id,
            electoral_term=parsed_vote.electoral_term,
            sitting_number=parsed_vote.sitting_number,
            vote_number=parsed_vote.vote_number,
            title=title,
            content_sha256=document.source.content_sha256,
            retrieved_at=datetime.fromisoformat(document.source.retrieved_at),
        )
        session.add(vote)
        session.flush()
        for row in parsed_vote.rows:
            session.add(NamedVoteRow(
                named_vote_id=vote.id,
                source_document_id=source.id,
                parliamentary_group=row.parliamentary_group,
                last_name=row.last_name,
                first_name=row.first_name,
                title=row.title,
                display_name=row.display_name,
                raw_outcome=row.raw_outcome,
                outcome=row.outcome,
                remark=row.remark,
                content_sha256=sha256(repr(row).encode()).hexdigest(),
                observed_at=datetime.fromisoformat(document.source.retrieved_at),
            ))
    session.commit()
    return parsed_vote


def _parse_row(row: dict[str, object]) -> VoteRow:
    active_outcomes = [header for header in OUTCOME_COLUMNS if _is_marked(row.get(header))]
    outcome = OUTCOME_COLUMNS[active_outcomes[0]] if len(active_outcomes) == 1 else "unknown"
    return VoteRow(
        parliamentary_group=_optional_string(row.get("fraktion/gruppe")),
        last_name=_optional_string(row.get("name")),
        first_name=_optional_string(row.get("vorname")),
        title=_optional_string(row.get("titel")),
        display_name=_optional_string(row.get("bezeichnung")),
        raw_outcome=active_outcomes[0] if len(active_outcomes) == 1 else ",".join(active_outcomes),
        outcome=outcome,
        remark=_optional_string(row.get("bemerkung")),
    )


def _get_or_create_source(session: Session, requested_url: str, document: FetchedDocument) -> SourceDocument:
    source = session.scalar(select(SourceDocument).where(SourceDocument.resolved_url == document.source.source_url, SourceDocument.content_sha256 == document.source.content_sha256))
    if source is not None:
        return source
    source = SourceDocument(
        publisher="Deutscher Bundestag",
        requested_url=requested_url,
        resolved_url=document.source.source_url,
        retrieved_at=datetime.fromisoformat(document.source.retrieved_at),
        status_code=document.source.status_code,
        content_type=document.source.content_type,
        content_sha256=document.source.content_sha256,
        retrieval_tool="PolitiklarCrawler/0.1",
        snapshot_location=document.snapshot_location,
    )
    session.add(source)
    session.flush()
    return source


def _normalize_header(value: object) -> str:
    return str(value).strip().lower() if value is not None else ""


def _required_int(row: dict[str, object], column: str) -> int:
    value = row.get(column)
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    raise ValueError(f"The XLSX file has no integer value for {column}.")


def _optional_string(value: object) -> str | None:
    return str(value).strip() if value is not None and str(value).strip() else None


def _is_marked(value: object) -> bool:
    return value in (1, 1.0, "1", "x", "X")